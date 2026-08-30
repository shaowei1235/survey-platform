from __future__ import annotations

from uuid import uuid4

from fastapi.testclient import TestClient
from sqlalchemy import select

from app.db import SessionLocal
from app.models import Response, User

from tests.conftest import LIKERT_ONE, PASSWORD, assert_error, auth_header, login


def test_qa01_wrong_password_same_as_unknown_user(api: TestClient) -> None:
    wrong = api.post("/api/v1/auth/login", json={"employee_no": "E-HR", "password": "wrong-pass-1"})
    unknown = api.post("/api/v1/auth/login", json={"employee_no": "E-NO-SUCH", "password": "wrong-pass-1"})
    assert_error(wrong, 401, "INVALID_CREDENTIALS")
    assert_error(unknown, 401, "INVALID_CREDENTIALS")
    assert wrong.json() == unknown.json()


def test_qa02_employee_cannot_use_admin_surveys(api: TestClient, sales_emp_auth: dict[str, str]) -> None:
    res = api.get("/api/v1/surveys", headers=sales_emp_auth)
    assert_error(res, 403, "FORBIDDEN_ROLE")


def test_qa03_refresh_token_rotates(api: TestClient) -> None:
    first = login(api, "E-HR")
    refresh = api.post("/api/v1/auth/refresh", json={"refresh_token": first["refresh_token"]})
    assert refresh.status_code == 200, refresh.text
    reused = api.post("/api/v1/auth/refresh", json={"refresh_token": first["refresh_token"]})
    assert_error(reused, 401, "UNAUTHENTICATED")


def test_qa04_create_user_without_department(api: TestClient, hr_auth: dict[str, str]) -> None:
    res = api.post(
        "/api/v1/users",
        headers=hr_auth,
        json={
            "employee_no": f"E-QA-{uuid4().hex[:8]}",
            "display_name": "欠部門",
            "password": "Init#pass1",
            "job_grade_id": "00000000-0000-0000-0000-000000000001",
            "generation": "20s",
            "roles": [{"role": "employee"}],
        },
    )
    assert_error(res, 400, "VALIDATION_ERROR")


def test_qa05_publish_empty_component_list(api: TestClient, hr_auth: dict[str, str]) -> None:
    created = api.post("/api/v1/surveys", headers=hr_auth, json={"title": f"QA empty {uuid4().hex[:8]}"})
    assert created.status_code == 201, created.text
    res = api.post(f"/api/v1/surveys/{created.json()['id']}/publish", headers=hr_auth)
    assert_error(res, 400, "VALIDATION_ERROR")


def test_qa06_patch_published_survey(api: TestClient, hr_auth: dict[str, str], demo_survey_id: str) -> None:
    res = api.patch(
        f"/api/v1/surveys/{demo_survey_id}",
        headers=hr_auth,
        json={"title": "should not change"},
    )
    assert_error(res, 409, "SURVEY_NOT_EDITABLE")


def test_qa09_submit_ignores_body_department(
    api: TestClient,
    tokens: dict[str, dict],
    hr_auth: dict[str, str],
    depts: dict[str, str],
) -> None:
    created = api.post("/api/v1/surveys", headers=hr_auth, json={"title": f"QA snap {uuid4().hex[:8]}"})
    survey_id = created.json()["id"]
    patched = api.patch(f"/api/v1/surveys/{survey_id}", headers=hr_auth, json={"component_list": LIKERT_ONE})
    assert patched.status_code == 200, patched.text
    published = api.post(f"/api/v1/surveys/{survey_id}/publish", headers=hr_auth)
    assert published.status_code == 200, published.text

    headers = auth_header(tokens, "E-DEV-02")
    res = api.post(
        f"/api/v1/client/surveys/{survey_id}/responses",
        headers=headers,
        json={
            "department_id": depts["営業部"],
            "answers": [{"fe_id": "c_q1", "type": "radio", "value": "5"}],
        },
    )
    assert res.status_code == 201, res.text

    db = SessionLocal()
    try:
        user = db.scalar(select(User).where(User.employee_no == "E-DEV-02"))
        row = db.scalar(select(Response).where(Response.survey_id == survey_id, Response.user_id == user.id))
        assert row is not None
        assert str(row.department_id) == depts["開発一課"]
        assert str(row.department_id) != depts["営業部"]
    finally:
        db.close()


def test_qa10_duplicate_submit(api: TestClient, sales_emp_auth: dict[str, str], demo_survey_id: str) -> None:
    res = api.post(
        f"/api/v1/client/surveys/{demo_survey_id}/responses",
        headers=sales_emp_auth,
        json={"answers": [{"fe_id": "c_q1", "type": "radio", "value": "4"}]},
    )
    assert_error(res, 409, "ALREADY_SUBMITTED")


def test_qa11_unpublished_get_and_post(
    api: TestClient, tokens: dict[str, dict], draft_survey_id: str
) -> None:
    headers = auth_header(tokens, "E-DEV-03")
    get_res = api.get(f"/api/v1/client/surveys/{draft_survey_id}", headers=headers)
    assert_error(get_res, 409, "NOT_PUBLISHED")
    post_res = api.post(
        f"/api/v1/client/surveys/{draft_survey_id}/responses",
        headers=headers,
        json={"answers": [{"fe_id": "c_q1", "type": "radio", "value": "4"}]},
    )
    assert_error(post_res, 409, "NOT_PUBLISHED")


def test_qa12_sales_manager_cannot_query_dev(
    api: TestClient, sales_mgr_auth: dict[str, str], demo_survey_id: str, depts: dict[str, str]
) -> None:
    res = api.get(
        "/api/v1/analytics/cross-tab",
        headers=sales_mgr_auth,
        params={"survey_id": demo_survey_id, "department_id": depts["開発部"]},
    )
    assert_error(res, 403, "FORBIDDEN_SCOPE")


def test_qa13_sales_manager_unfiltered_rows(
    api: TestClient, sales_mgr_auth: dict[str, str], demo_survey_id: str, depts: dict[str, str]
) -> None:
    res = api.get("/api/v1/analytics/cross-tab", headers=sales_mgr_auth, params={"survey_id": demo_survey_id})
    assert res.status_code == 200, res.text
    names = {row["department_name"] for row in res.json()["rows"]}
    assert names == {"営業部", "第一営業課", "第二営業課"}
    by_name = {row["department_name"]: row["department_id"] for row in res.json()["rows"]}
    assert by_name["営業部"] == depts["営業部"]
    assert by_name["第一営業課"] == depts["第一営業課"]
    assert by_name["第二営業課"] == depts["第二営業課"]


def test_qa14_employee_cannot_run_intent(
    api: TestClient, sales_emp_auth: dict[str, str], demo_survey_id: str, depts: dict[str, str]
) -> None:
    res = api.post(
        "/api/v1/analytics/intent",
        headers=sales_emp_auth,
        json={"survey_id": demo_survey_id, "intent": "dept_low_score_and_causes", "department_id": depts["営業部"]},
    )
    assert_error(res, 403, "FORBIDDEN_ROLE")


def test_qa15_dept_manager_masked_hides_n(
    api: TestClient, dev_mgr_auth: dict[str, str], demo_survey_id: str, depts: dict[str, str]
) -> None:
    res = api.get(
        "/api/v1/analytics/cross-tab",
        headers=dev_mgr_auth,
        params={"survey_id": demo_survey_id, "department_id": depts["開発部"]},
    )
    assert res.status_code == 200, res.text
    cells = res.json()["rows"][0]["cells"]
    assert cells
    assert all(c["masked"] is True for c in cells)
    assert all(c["avg_score"] is None for c in cells)
    assert all(c["n"] is None for c in cells)


def test_qa16_hr_sees_n_on_masked_hr_dept(
    api: TestClient, hr_auth: dict[str, str], demo_survey_id: str, depts: dict[str, str]
) -> None:
    res = api.get(
        "/api/v1/analytics/cross-tab",
        headers=hr_auth,
        params={"survey_id": demo_survey_id, "department_id": depts["人事部"]},
    )
    assert res.status_code == 200, res.text
    cells = res.json()["rows"][0]["cells"]
    assert cells
    assert all(c["masked"] is True for c in cells)
    assert all(c["avg_score"] is None for c in cells)
    assert all(c["n"] == 3 for c in cells)


def test_qa_sales_scores_visible(api: TestClient, hr_auth: dict[str, str], demo_survey_id: str, depts: dict[str, str]) -> None:
    res = api.get(
        "/api/v1/analytics/cross-tab",
        headers=hr_auth,
        params={"survey_id": demo_survey_id, "department_id": depts["営業部"]},
    )
    assert res.status_code == 200, res.text
    cells = res.json()["rows"][0]["cells"]
    assert [c["avg_score"] for c in cells] == [4.25, 4.375, 3.0, 3.5, 4.375]
    assert all(c["masked"] is False and c["n"] == 8 for c in cells)


def test_qa_sales_ka1_scores_visible(api: TestClient, hr_auth: dict[str, str], demo_survey_id: str, depts: dict[str, str]) -> None:
    res = api.get(
        "/api/v1/analytics/cross-tab",
        headers=hr_auth,
        params={"survey_id": demo_survey_id, "department_id": depts["第一営業課"]},
    )
    assert res.status_code == 200, res.text
    cells = res.json()["rows"][0]["cells"]
    assert [c["avg_score"] for c in cells] == [4.2, 4.4, 3.2, 3.6, 4.4]
    assert all(c["masked"] is False and c["n"] == 5 for c in cells)


def test_qa_sales_ka2_masked(api: TestClient, hr_auth: dict[str, str], demo_survey_id: str, depts: dict[str, str]) -> None:
    res = api.get(
        "/api/v1/analytics/cross-tab",
        headers=hr_auth,
        params={"survey_id": demo_survey_id, "department_id": depts["第二営業課"]},
    )
    assert res.status_code == 200, res.text
    cells = res.json()["rows"][0]["cells"]
    assert cells
    assert all(c["masked"] is True for c in cells)
    assert all(c["avg_score"] is None for c in cells)
    assert all(c["n"] == 3 for c in cells)


def test_qa17_empty_survey_cross_tab_no_llm(api: TestClient, hr_auth: dict[str, str]) -> None:
    created = api.post("/api/v1/surveys", headers=hr_auth, json={"title": f"QA empty-ans {uuid4().hex[:8]}"})
    survey_id = created.json()["id"]
    api.patch(f"/api/v1/surveys/{survey_id}", headers=hr_auth, json={"component_list": LIKERT_ONE})
    published = api.post(f"/api/v1/surveys/{survey_id}/publish", headers=hr_auth)
    assert published.status_code == 200, published.text
    res = api.get("/api/v1/analytics/cross-tab", headers=hr_auth, params={"survey_id": survey_id})
    assert res.status_code == 200, res.text
    rows = res.json()["rows"]
    assert rows
    assert all(all(c["masked"] is True and c["avg_score"] is None for c in row["cells"]) for row in rows)


def test_qa19_intent_without_llm(
    api: TestClient, hr_auth: dict[str, str], demo_survey_id: str, depts: dict[str, str], monkeypatch
) -> None:
    monkeypatch.setattr("app.services.llm.settings.llm_base_url", "")
    monkeypatch.setattr("app.services.llm.settings.llm_api_key", "")
    monkeypatch.setattr("app.services.llm.settings.llm_model", "")
    res = api.post(
        "/api/v1/analytics/intent",
        headers=hr_auth,
        json={"survey_id": demo_survey_id, "intent": "dept_low_score_and_causes", "department_id": depts["営業部"]},
    )
    assert_error(res, 502, "AI_UPSTREAM_FAILED")
    assert "conclusion" not in res.json()


def test_qa18_intent_mock_numbers_match_table(
    api: TestClient, hr_auth: dict[str, str], demo_survey_id: str, depts: dict[str, str], monkeypatch
) -> None:
    monkeypatch.setattr(
        "app.routers.analytics.complete_json",
        lambda *_args, **_kwargs: {"conclusion": "営業部では労働時間と人事制度への不満が相対的に目立ちます。"},
    )
    tab = api.get(
        "/api/v1/analytics/cross-tab",
        headers=hr_auth,
        params={"survey_id": demo_survey_id, "department_id": depts["営業部"]},
    )
    assert tab.status_code == 200, tab.text
    res = api.post(
        "/api/v1/analytics/intent",
        headers=hr_auth,
        json={"survey_id": demo_survey_id, "intent": "dept_low_score_and_causes", "department_id": depts["営業部"]},
    )
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["conclusion"]
    lows = body["evidence"]["low_questions"]
    assert lows
    by_id = {c["fe_id"]: c["avg_score"] for c in tab.json()["rows"][0]["cells"] if not c["masked"]}
    for item in lows:
        assert item["avg_score"] == by_id[item["fe_id"]]
    assert all("E-SALES-01" not in (q.get("text") or "") for q in body["evidence"].get("quotes") or [])
    texts = [q.get("text") or "" for q in body["evidence"].get("quotes") or []]
    assert any("月末の残業" in text for text in texts)
    assert any("少人数のため" in text for text in texts)


def test_qa20_intent_mock_invented_score(
    api: TestClient, hr_auth: dict[str, str], demo_survey_id: str, depts: dict[str, str], monkeypatch
) -> None:
    monkeypatch.setattr(
        "app.routers.analytics.complete_json",
        lambda *_args, **_kwargs: {"conclusion": "平均点は1.11です。"},
    )
    res = api.post(
        "/api/v1/analytics/intent",
        headers=hr_auth,
        json={"survey_id": demo_survey_id, "intent": "dept_low_score_and_causes", "department_id": depts["営業部"]},
    )
    assert_error(res, 502, "AI_NUMBER_MISMATCH")


def test_qa24_closed_survey_cannot_be_answered(api: TestClient, tokens: dict[str, dict], hr_auth: dict[str, str]) -> None:
    created = api.post("/api/v1/surveys", headers=hr_auth, json={"title": f"QA close {uuid4().hex[:8]}"})
    survey_id = created.json()["id"]
    api.patch(f"/api/v1/surveys/{survey_id}", headers=hr_auth, json={"component_list": LIKERT_ONE})
    assert api.post(f"/api/v1/surveys/{survey_id}/publish", headers=hr_auth).status_code == 200
    assert api.post(f"/api/v1/surveys/{survey_id}/close", headers=hr_auth).status_code == 200
    headers = auth_header(tokens, "E-DEV-03")
    get_res = api.get(f"/api/v1/client/surveys/{survey_id}", headers=headers)
    assert_error(get_res, 409, "NOT_PUBLISHED")
    post_res = api.post(
        f"/api/v1/client/surveys/{survey_id}/responses",
        headers=headers,
        json={"answers": [{"fe_id": "c_q1", "type": "radio", "value": "5"}]},
    )
    assert_error(post_res, 409, "NOT_PUBLISHED")


def test_dept_manager_can_list_job_grades(api: TestClient, sales_mgr_auth: dict[str, str]) -> None:
    res = api.get("/api/v1/job-grades", headers=sales_mgr_auth)
    assert res.status_code == 200, res.text
    names = {g["name"] for g in res.json()["items"]}
    assert {"一般", "部長"} <= names


def test_employee_cannot_list_job_grades(api: TestClient, sales_emp_auth: dict[str, str]) -> None:
    res = api.get("/api/v1/job-grades", headers=sales_emp_auth)
    assert_error(res, 403, "FORBIDDEN_ROLE")


def test_cross_tab_generation_filter(
    api: TestClient, hr_auth: dict[str, str], demo_survey_id: str, depts: dict[str, str]
) -> None:
    visible = api.get(
        "/api/v1/analytics/cross-tab",
        headers=hr_auth,
        params={"survey_id": demo_survey_id, "department_id": depts["営業部"], "generation": "20s"},
    )
    assert visible.status_code == 200, visible.text
    assert [c["avg_score"] for c in visible.json()["rows"][0]["cells"]] == [4.25, 4.375, 3.0, 3.5, 4.375]
    hidden = api.get(
        "/api/v1/analytics/cross-tab",
        headers=hr_auth,
        params={"survey_id": demo_survey_id, "department_id": depts["営業部"], "generation": "50s"},
    )
    assert hidden.status_code == 200, hidden.text
    cells = hidden.json()["rows"][0]["cells"]
    assert all(c["masked"] is True and c["avg_score"] is None for c in cells)


def test_cross_tab_job_grade_filter(
    api: TestClient, hr_auth: dict[str, str], demo_survey_id: str, depts: dict[str, str]
) -> None:
    grades = {g["name"]: g["id"] for g in api.get("/api/v1/job-grades", headers=hr_auth).json()["items"]}
    visible = api.get(
        "/api/v1/analytics/cross-tab",
        headers=hr_auth,
        params={"survey_id": demo_survey_id, "department_id": depts["営業部"], "job_grade_id": grades["一般"]},
    )
    assert visible.status_code == 200, visible.text
    assert [c["avg_score"] for c in visible.json()["rows"][0]["cells"]] == [4.25, 4.375, 3.0, 3.5, 4.375]
    hidden = api.get(
        "/api/v1/analytics/cross-tab",
        headers=hr_auth,
        params={"survey_id": demo_survey_id, "department_id": depts["営業部"], "job_grade_id": grades["部長"]},
    )
    assert hidden.status_code == 200, hidden.text
    cells = hidden.json()["rows"][0]["cells"]
    assert all(c["masked"] is True and c["avg_score"] is None for c in cells)


def _xlsx_data_rows(content: bytes) -> dict[str, list]:
    from io import BytesIO

    from openpyxl import load_workbook

    sheet = load_workbook(BytesIO(content), data_only=True).active
    header_row = next(row[0].row for row in sheet.iter_rows(min_row=1, max_col=1) if row[0].value == "部門")
    question_count = sum(1 for cell in sheet[header_row][1:] if cell.value)
    out: dict[str, list] = {}
    for row in sheet.iter_rows(min_row=header_row + 1):
        name = row[0].value
        if not name:
            continue
        values = []
        for cell in row[1 : 1 + question_count]:
            value = cell.value
            if isinstance(value, (int, float)):
                values.append(round(float(value), 1))
            else:
                values.append(value)
        out[str(name)] = values
    return out


def test_export_sales_matches_screen_and_masks_ka2(
    api: TestClient, hr_auth: dict[str, str], demo_survey_id: str
) -> None:
    from app.services.analytics import MASKED_CELL_LABEL

    res = api.get("/api/v1/analytics/cross-tab.xlsx", headers=hr_auth, params={"survey_id": demo_survey_id})
    assert res.status_code == 200, res.text
    assert "spreadsheetml" in res.headers.get("content-type", "")
    assert "filename*=UTF-8''" in res.headers.get("content-disposition", "")
    rows = _xlsx_data_rows(res.content)
    assert rows["営業部"] == [4.3, 4.4, 3.0, 3.5, 4.4]
    assert rows["第一営業課"] == [4.2, 4.4, 3.2, 3.6, 4.4]
    assert rows["第二営業課"] == [MASKED_CELL_LABEL] * 5
    assert b"4.3333" not in res.content
    assert b"2.3333" not in res.content


def test_export_hr_dept_has_no_scores(
    api: TestClient, hr_auth: dict[str, str], demo_survey_id: str, depts: dict[str, str]
) -> None:
    from app.services.analytics import MASKED_CELL_LABEL

    res = api.get(
        "/api/v1/analytics/cross-tab.xlsx",
        headers=hr_auth,
        params={"survey_id": demo_survey_id, "department_id": depts["人事部"]},
    )
    assert res.status_code == 200, res.text
    rows = _xlsx_data_rows(res.content)
    assert list(rows) == ["人事部"]
    assert rows["人事部"] == [MASKED_CELL_LABEL] * 5
    assert b"2.3333" not in res.content


def test_export_sales_manager_cannot_export_dev(
    api: TestClient, sales_mgr_auth: dict[str, str], demo_survey_id: str, depts: dict[str, str]
) -> None:
    res = api.get(
        "/api/v1/analytics/cross-tab.xlsx",
        headers=sales_mgr_auth,
        params={"survey_id": demo_survey_id, "department_id": depts["開発部"]},
    )
    assert_error(res, 403, "FORBIDDEN_SCOPE")


def test_employee_cannot_export(api: TestClient, sales_emp_auth: dict[str, str], demo_survey_id: str) -> None:
    res = api.get("/api/v1/analytics/cross-tab.xlsx", headers=sales_emp_auth, params={"survey_id": demo_survey_id})
    assert_error(res, 403, "FORBIDDEN_ROLE")

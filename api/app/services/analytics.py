from __future__ import annotations

from collections import defaultdict
from decimal import Decimal, ROUND_HALF_UP
from io import BytesIO
from statistics import mean
from urllib.parse import quote
from uuid import UUID

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.config import settings
from app.errors import ApiError
from app.models import AnswerItem, Department, Response, Survey, SurveyStatus, User, UserGeneration, UserRole, UserRoleCode
from app.services.authz import ANALYST_ROLES, assert_department_in_scope, descendant_ids, is_dept_manager_only, require_any_role, scope_department_ids
from app.services.components import likert_questions, option_score
from app.services.deidentify import sanitize_quote

MASKED_CELL_LABEL = "非開示（n<5）"
_GEN_LABELS = {"20s": "20代", "30s": "30代", "40s": "40代", "50s": "50代", "60s_plus": "60代以上"}
_ALL_LABEL = "すべて"


def subtree_department_ids(db: Session, user: User, department_id: UUID) -> list[UUID]:
    return descendant_ids(db, user.company_id, [department_id])


def _mask_roll_up(n: int, min_n: int, own_n: int, child_ns: list[int]) -> bool:
    if n < min_n:
        return True
    positive = [count for count in child_ns if count > 0]
    if own_n == 0 and len(positive) == 1 and positive[0] < min_n:
        return True
    return False


def _leaf_rows(rows: list[dict], depts: dict[UUID, Department]) -> list[dict]:
    present = {r["department_id"] for r in rows}
    children: dict[UUID, list[UUID]] = {}
    for dept in depts.values():
        if dept.parent_id:
            children.setdefault(dept.parent_id, []).append(dept.id)

    def has_present_descendant(dept_id: UUID) -> bool:
        stack = list(children.get(dept_id, []))
        while stack:
            current = stack.pop()
            if current in present:
                return True
            stack.extend(children.get(current, []))
        return False

    return [row for row in rows if not has_present_descendant(row["department_id"])]


def _survey_for_company(db: Session, user: User, survey_id: UUID) -> Survey:
    survey = db.scalar(select(Survey).where(Survey.id == survey_id, Survey.company_id == user.company_id))
    if survey is None:
        raise ApiError("NOT_FOUND")
    if survey.status == SurveyStatus.draft:
        raise ApiError("NOT_FOUND")
    return survey


def _response_query(db: Session, survey_id: UUID, dept_ids: list[UUID], generation: str | None, job_grade_id: UUID | None):
    stmt = select(Response).where(Response.survey_id == survey_id, Response.department_id.in_(dept_ids))
    if generation:
        stmt = stmt.where(Response.generation == UserGeneration(generation))
    if job_grade_id:
        stmt = stmt.where(Response.job_grade_id == job_grade_id)
    return db.scalars(stmt).all()


def cross_tab(
    db: Session,
    user: User,
    survey_id: UUID,
    department_id: UUID | None,
    generation: str | None,
    job_grade_id: UUID | None,
) -> dict:
    require_any_role(user, ANALYST_ROLES)
    survey = _survey_for_company(db, user, survey_id)
    scope = scope_department_ids(db, user)
    if department_id:
        assert_department_in_scope(db, user, department_id)
        row_depts = [department_id]
    else:
        row_depts = scope
    all_depts = {
        d.id: d for d in db.scalars(select(Department).where(Department.company_id == user.company_id)).all()
    }
    row_depts = sorted(
        row_depts,
        key=lambda did: (all_depts[did].sort_order, all_depts[did].name) if did in all_depts else (999, ""),
    )
    subtrees = {did: descendant_ids(db, user.company_id, [did]) for did in row_depts}
    member_ids = list({mid for ids in subtrees.values() for mid in ids})
    questions = likert_questions(survey.component_list or [])
    responses = _response_query(db, survey.id, member_ids, generation, job_grade_id)
    resp_ids = [r.id for r in responses]
    items = []
    if resp_ids:
        items = db.scalars(select(AnswerItem).where(AnswerItem.response_id.in_(resp_ids))).all()
    by_resp_dept = {r.id: r.department_id for r in responses}
    scores: dict[tuple[UUID, str], list[int]] = defaultdict(list)
    for item in items:
        dept_id = by_resp_dept.get(item.response_id)
        if dept_id is None:
            continue
        q = next((x for x in questions if x["fe_id"] == item.fe_id), None)
        if q is None:
            continue
        scored = option_score(q, str(item.value))
        if scored is None:
            continue
        scores[(dept_id, item.fe_id)].append(scored)
    hide_n = is_dept_manager_only(user)
    min_n = settings.min_cell_n
    rows = []
    for dept_id in row_depts:
        dept = all_depts.get(dept_id)
        if dept is None:
            continue
        member = subtrees[dept_id]
        child_ids = [mid for mid in member if mid != dept_id]
        cells = []
        for q in questions:
            values: list[int] = []
            for mid in member:
                values.extend(scores.get((mid, q["fe_id"]), []))
            own_n = len(scores.get((dept_id, q["fe_id"]), []))
            child_ns = [len(scores.get((cid, q["fe_id"]), [])) for cid in child_ids]
            n = len(values)
            masked = _mask_roll_up(n, min_n, own_n, child_ns)
            cell = {
                "fe_id": q["fe_id"],
                "n": None if (masked and hide_n) else n,
                "avg_score": None if masked else (round(mean(values), 4) if values else None),
                "masked": masked,
            }
            cells.append(cell)
        rows.append({"department_id": dept_id, "department_name": dept.name, "cells": cells})
    return {
        "survey_id": survey.id,
        "questions": [
            {"fe_id": q["fe_id"], "title": (q.get("props") or {}).get("title", ""), "score_enabled": True} for q in questions
        ],
        "rows": rows,
        "charts": [{"id": "bar_dept_scores", "type": "bar", "title_key": "chart.dept_scores"}],
        "empty": not any(scores.values()),
    }


def completion(db: Session, user: User, survey_id: UUID, department_id: UUID | None) -> dict:
    require_any_role(user, ANALYST_ROLES)
    survey = _survey_for_company(db, user, survey_id)
    scope = scope_department_ids(db, user)
    if department_id:
        assert_department_in_scope(db, user, department_id)
        row_depts = [department_id]
    else:
        row_depts = scope
    all_depts = {
        d.id: d for d in db.scalars(select(Department).where(Department.company_id == user.company_id)).all()
    }
    row_depts = sorted(
        row_depts,
        key=lambda did: (all_depts[did].sort_order, all_depts[did].name) if did in all_depts else (999, ""),
    )
    subtrees = {did: descendant_ids(db, user.company_id, [did]) for did in row_depts}
    employee_ids = select(UserRole.user_id).where(UserRole.role == UserRoleCode.employee)
    people = db.scalars(
        select(User).where(
            User.company_id == user.company_id,
            User.is_active.is_(True),
            User.id.in_(employee_ids),
        )
    ).all()
    submitted_ids = set(db.scalars(select(Response.user_id).where(Response.survey_id == survey.id)).all())
    eligible_by: dict[UUID, int] = defaultdict(int)
    submitted_by: dict[UUID, int] = defaultdict(int)
    for person in people:
        eligible_by[person.department_id] += 1
        if person.id in submitted_ids:
            submitted_by[person.department_id] += 1
    hide_n = is_dept_manager_only(user)
    min_n = settings.min_cell_n
    rows = []
    for dept_id in row_depts:
        dept = all_depts.get(dept_id)
        if dept is None:
            continue
        member = subtrees[dept_id]
        child_ids = [mid for mid in member if mid != dept_id]
        eligible = sum(eligible_by[mid] for mid in member)
        submitted = sum(submitted_by[mid] for mid in member)
        unanswered = eligible - submitted
        own_n = eligible_by[dept_id]
        child_ns = [eligible_by[cid] for cid in child_ids]
        masked = _mask_roll_up(eligible, min_n, own_n, child_ns)
        rate = round(submitted / eligible, 4) if eligible and not masked else None
        rows.append(
            {
                "department_id": dept_id,
                "department_name": dept.name,
                "eligible": None if (masked and hide_n) else eligible,
                "submitted": None if masked else submitted,
                "unanswered": None if masked else unanswered,
                "rate": rate,
                "masked": masked,
            }
        )
    return {"survey_id": survey.id, "rows": rows}


def build_intent_evidence(db: Session, user: User, survey: Survey, department_id: UUID) -> dict:
    tab = cross_tab(db, user, survey.id, department_id, None, None)
    company_tab = cross_tab(db, user, survey.id, None, None, None)
    row = next((r for r in tab["rows"] if r["department_id"] == department_id), None)
    if row is None:
        raise ApiError("NOT_FOUND")
    low = []
    for cell, q in zip(row["cells"], tab["questions"], strict=False):
        if cell["masked"] or cell["avg_score"] is None:
            continue
        low.append(
            {
                "fe_id": cell["fe_id"],
                "title": q["title"],
                "avg_score": cell["avg_score"],
                "n": cell["n"],
                "masked": False,
            }
        )
    low.sort(key=lambda x: x["avg_score"])
    low = low[:5]
    nos = list(db.scalars(select(User.employee_no).where(User.company_id == user.company_id)))
    quotes = collect_quotes(db, survey, subtree_department_ids(db, user, department_id), nos)
    dept = db.scalar(select(Department).where(Department.id == department_id))
    dept_avgs = {c["fe_id"]: c["avg_score"] for c in row["cells"] if not c["masked"]}
    all_depts = {
        d.id: d for d in db.scalars(select(Department).where(Department.company_id == user.company_id)).all()
    }
    bench_rows = _leaf_rows(company_tab["rows"], all_depts)
    bench: dict[str, float] = {}
    for q in tab["questions"]:
        vals = []
        for r in bench_rows:
            cell = next(
                (c for c in r["cells"] if c["fe_id"] == q["fe_id"] and not c["masked"] and c["avg_score"] is not None),
                None,
            )
            if cell:
                vals.append(cell["avg_score"])
        if vals:
            bench[q["fe_id"]] = round(sum(vals) / len(vals), 4)
    evidence = {
        "department_id": str(department_id),
        "department_name": dept.name if dept else "",
        "department_avg_by_fe_id": dept_avgs,
        "low_questions": low,
        "quotes": quotes,
        "charts": [{"id": "bar_dept_scores", "type": "bar", "title_key": "chart.dept_scores"}],
        "empty": bool(tab.get("empty")),
    }
    if is_dept_manager_only(user):
        evidence["benchmark_avg_by_fe_id"] = bench
    else:
        evidence["company_avg_by_fe_id"] = bench
        evidence["benchmark_avg_by_fe_id"] = bench
    return evidence


def collect_quotes(db: Session, survey: Survey, dept_ids: list[UUID], employee_nos: list[str], fe_ids: set[str] | None = None) -> list[dict]:
    responses = db.scalars(
        select(Response).where(Response.survey_id == survey.id, Response.department_id.in_(dept_ids))
    ).all()
    if not responses:
        return []
    items = db.scalars(
        select(AnswerItem).where(AnswerItem.response_id.in_([r.id for r in responses]), AnswerItem.type == "textarea")
    ).all()
    quotes = []
    for item in items:
        if fe_ids is not None and item.fe_id not in fe_ids:
            continue
        text = item.value if isinstance(item.value, str) else None
        cleaned = sanitize_quote(text or "", employee_nos)
        if cleaned:
            quotes.append({"text": cleaned, "fe_id": item.fe_id})
        if len(quotes) >= 30:
            break
    return quotes


def format_score(value: float) -> Decimal:
    return Decimal(str(value)).quantize(Decimal("0.1"), rounding=ROUND_HALF_UP)


def _safe_filename(title: str) -> str:
    cleaned = "".join("_" if ch in '\\/:*?"<>|' else ch for ch in title).strip()
    cleaned = " ".join(cleaned.split())[:40].rstrip(" .")
    return cleaned or "集計"


def cross_tab_xlsx_filename(survey_title: str) -> str:
    return f"集計_{_safe_filename(survey_title)}.xlsx"


def content_disposition(filename: str) -> str:
    ascii_name = "crosstab.xlsx"
    return f"attachment; filename=\"{ascii_name}\"; filename*=UTF-8''{quote(filename)}"


def build_cross_tab_xlsx(
    tab: dict,
    *,
    survey_title: str,
    department_name: str | None,
    generation: str | None,
    job_grade_name: str | None,
) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = "集計"
    meta = [
        ("アンケート", survey_title),
        ("対象部門", department_name or _ALL_LABEL),
        ("年代", _GEN_LABELS.get(generation or "", _ALL_LABEL) if generation else _ALL_LABEL),
        ("職級", job_grade_name or _ALL_LABEL),
        ("注記", "回答数が5未満のセルは非開示です。画面と同じ集計です。"),
    ]
    label_font = Font(bold=True)
    for index, (label, value) in enumerate(meta, start=1):
        sheet.cell(index, 1, label).font = label_font
        sheet.cell(index, 2, value)
    header_row = 7
    sheet.cell(header_row, 1, "部門").font = label_font
    questions = tab.get("questions") or []
    for col, question in enumerate(questions, start=2):
        cell = sheet.cell(header_row, col, question.get("title") or question.get("fe_id") or "")
        cell.font = label_font
        cell.alignment = Alignment(wrap_text=True)
    for row_index, row in enumerate(tab.get("rows") or [], start=header_row + 1):
        sheet.cell(row_index, 1, row.get("department_name") or "")
        by_fe = {cell["fe_id"]: cell for cell in row.get("cells") or []}
        for col, question in enumerate(questions, start=2):
            item = by_fe.get(question["fe_id"])
            dest = sheet.cell(row_index, col)
            if item is None or item.get("masked") or item.get("avg_score") is None:
                dest.value = MASKED_CELL_LABEL
            else:
                dest.value = float(format_score(item["avg_score"]))
                dest.number_format = "0.0"
    sheet.column_dimensions["A"].width = 18
    for col in range(2, 2 + max(len(questions), 1)):
        sheet.column_dimensions[sheet.cell(header_row, col).column_letter].width = 18
    sheet.freeze_panes = "B8"
    buf = BytesIO()
    book.save(buf)
    return buf.getvalue()
